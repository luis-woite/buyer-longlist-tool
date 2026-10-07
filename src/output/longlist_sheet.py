"""
Longlist Excel output.

Writes the researched buyer longlist to a formatted workbook — one
sheet for strategic buyers, one for financial buyers, each with the
grounding evidence visible right next to the rationale, not hidden in
a tooltip or a separate document.
"""

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ACCENT_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=14, color="1F3864")
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="666666")

COLUMNS = [
    ("Buyer Name", "name", 26),
    ("Rationale", "rationale", 50),
    ("Evidence / Source", "evidence_source", 40),
]


def _write_buyer_sheet(wb: Workbook, sheet_name: str, title: str, buyers: list[dict]):
    ws = wb.create_sheet(sheet_name)

    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))

    header_row = 3
    for col_idx, (header, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = ACCENT_FILL
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    if not buyers:
        ws.cell(row=header_row + 1, column=1, value="No candidates found meeting the grounding bar for this category.").font = NOTE_FONT
    else:
        for i, buyer in enumerate(buyers, start=1):
            r = header_row + i
            for col_idx, (_, key, _) in enumerate(COLUMNS, start=1):
                cell = ws.cell(row=r, column=col_idx, value=buyer.get(key))
                cell.font = BODY_FONT
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = f"A{header_row + 1}"


def write_longlist(result: dict, output_path: str = "output/buyer_longlist.xlsx") -> str:
    """
    Writes a fresh longlist workbook for this research run. Returns the
    path written to. Each run overwrites the file — this is a point-in-time
    research output, not an accumulating tracker.
    """
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)  # drop the default sheet, we create our own named ones

    target_name = result.get("target_company", "Target")

    _write_buyer_sheet(
        wb, "Strategic Buyers", f"Strategic Buyer Longlist — {target_name}", result.get("strategic_buyers") or []
    )
    _write_buyer_sheet(
        wb, "Financial Buyers", f"Financial Buyer Longlist — {target_name}", result.get("financial_buyers") or []
    )

    # A small summary/notes sheet, placed first
    notes_ws = wb.create_sheet("Summary", 0)
    notes_ws.cell(row=1, column=1, value=f"Buyer Longlist — {target_name}").font = TITLE_FONT
    notes_ws.cell(row=2, column=1, value=f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = NOTE_FONT
    notes_ws.cell(row=4, column=1, value=f"Strategic buyers found: {len(result.get('strategic_buyers') or [])}").font = BODY_FONT
    notes_ws.cell(row=5, column=1, value=f"Financial buyers found: {len(result.get('financial_buyers') or [])}").font = BODY_FONT
    if result.get("research_notes"):
        notes_ws.cell(row=7, column=1, value="Research notes:").font = Font(name="Arial", bold=True, size=10)
        notes_cell = notes_ws.cell(row=8, column=1, value=result["research_notes"])
        notes_cell.font = BODY_FONT
        notes_cell.alignment = Alignment(wrap_text=True, vertical="top")
        notes_ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=3)
        notes_ws.row_dimensions[8].height = 45
    disclaimer = notes_ws.cell(
        row=10,
        column=1,
        value=(
            "This longlist was generated from web search via an LLM and is a first-pass research aid "
            "only — every candidate and rationale should be independently verified before use in an "
            "actual process."
        ),
    )
    disclaimer.font = NOTE_FONT
    disclaimer.alignment = Alignment(wrap_text=True)
    notes_ws.merge_cells(start_row=10, start_column=1, end_row=10, end_column=3)
    notes_ws.row_dimensions[10].height = 45
    for col, width in zip(["A", "B", "C"], [30, 30, 30]):
        notes_ws.column_dimensions[col].width = width

    wb.save(output_path)
    return output_path
